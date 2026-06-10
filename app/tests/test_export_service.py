import json
from collections.abc import Generator
from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.main import app
from app.models.amazon_data import AmazonProductData
from app.models.analysis import ProductAnalysis
from app.models.email_campaign import EmailCampaign
from app.models.product import Product
from app.models.supplier import Supplier
from app.services.export_service import PRODUCT_EXPORT_COLUMNS, SUPPLIER_EXPORT_COLUMNS, ExportService


@pytest.fixture()
def db_session() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    with TestingSessionLocal() as db:
        seed_export_data(db)
        yield db

    Base.metadata.drop_all(bind=engine)


@pytest.fixture()
def client(monkeypatch: pytest.MonkeyPatch) -> Generator[TestClient, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    with TestingSessionLocal() as db:
        seed_export_data(db)

    def override_get_db() -> Generator[Session, None, None]:
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(bind=engine)


def seed_export_data(db: Session) -> None:
    supplier_one = Supplier(
        supplier_name="Acme Supply",
        company="Acme Inc",
        email="sales@acme.test",
        website="https://acme.test",
        phone="555-0100",
        city="Miami",
        country="US",
        category="Toys",
        status="active",
        is_valid_email=True,
        notes="Preferred",
    )
    supplier_two = Supplier(supplier_name="Beta Supply", category="Kitchen", status="pending")
    db.add_all([supplier_one, supplier_two])
    db.flush()

    products = [
        Product(
            supplier_id=supplier_one.id,
            product_name="Buy Widget",
            sku="BUY-1",
            upc="111",
            brand="Acme",
            category="Toys",
            supplier_cost=Decimal("10.00"),
            status="buy",
        ),
        Product(
            supplier_id=supplier_one.id,
            product_name="Review Widget",
            sku="REV-1",
            upc="222",
            brand="Acme",
            category="Toys",
            supplier_cost=Decimal("12.00"),
            status="review",
        ),
        Product(
            supplier_id=supplier_two.id,
            product_name="Discard Gadget",
            sku="DIS-1",
            upc="333",
            brand="Beta",
            category="Kitchen",
            supplier_cost=Decimal("20.00"),
            status="discard",
        ),
    ]
    db.add_all(products)
    db.flush()

    db.add_all(
        [
            AmazonProductData(
                product_id=products[0].id,
                asin="BUYASIN",
                amazon_title="Amazon Buy Widget",
                amazon_url="https://amazon.test/buy",
                current_price=Decimal("25.00"),
                buybox_price=Decimal("24.50"),
                amazon_price=Decimal("25.00"),
                rating=Decimal("4.50"),
                reviews_count=100,
                sellers_count=4,
                estimated_sales=1000,
            ),
            AmazonProductData(
                product_id=products[1].id,
                asin="REVASIN",
                amazon_title="Amazon Review Widget",
                current_price=Decimal("22.00"),
                rating=Decimal("4.10"),
                reviews_count=50,
                sellers_count=6,
                estimated_sales=500,
            ),
            AmazonProductData(
                product_id=products[2].id,
                asin="DISASIN",
                amazon_title="Amazon Discard Gadget",
                current_price=Decimal("18.00"),
                sellers_count=12,
                estimated_sales=100,
            ),
        ]
    )
    db.add_all(
        [
            ProductAnalysis(
                product_id=products[0].id,
                net_profit=Decimal("5.00"),
                margin=Decimal("0.2000"),
                roi=Decimal("0.5000"),
                profitability_score=Decimal("70.00"),
                roi_score=Decimal("85.00"),
                sales_score=Decimal("100.00"),
                price_stability_score=Decimal("100.00"),
                sellers_score=Decimal("100.00"),
                data_quality_score=Decimal("100.00"),
                final_opportunity_score=Decimal("88.00"),
                recommendation_status="buy",
                recommendation_reason="Strong opportunity.",
                risks_json=json.dumps([]),
            ),
            ProductAnalysis(
                product_id=products[1].id,
                net_profit=Decimal("3.00"),
                margin=Decimal("0.1364"),
                roi=Decimal("0.2500"),
                profitability_score=Decimal("35.00"),
                roi_score=Decimal("50.00"),
                sales_score=Decimal("70.00"),
                price_stability_score=Decimal("80.00"),
                sellers_score=Decimal("80.00"),
                data_quality_score=Decimal("100.00"),
                final_opportunity_score=Decimal("59.50"),
                recommendation_status="review",
                recommendation_reason="Requires review.",
                risks_json=json.dumps(["Low ROI."]),
            ),
            ProductAnalysis(
                product_id=products[2].id,
                net_profit=Decimal("-8.00"),
                margin=Decimal("-0.4444"),
                roi=Decimal("-0.4000"),
                profitability_score=Decimal("0.00"),
                roi_score=Decimal("10.00"),
                sales_score=Decimal("25.00"),
                price_stability_score=Decimal("40.00"),
                sellers_score=Decimal("35.00"),
                data_quality_score=Decimal("90.00"),
                final_opportunity_score=Decimal("21.00"),
                recommendation_status="discard",
                recommendation_reason="Negative profit.",
                risks_json=json.dumps(["Negative net profit."]),
            ),
        ]
    )
    db.add(
        EmailCampaign(
            subject="Supplier outreach",
            template_id="template-1",
            status="completed",
            total_recipients=2,
            sent_count=2,
            failed_count=0,
            created_at=datetime.utcnow(),
            updated_at=datetime.utcnow(),
        )
    )
    db.commit()


def workbook_from_buffer(buffer: BytesIO):
    buffer.seek(0)
    return load_workbook(buffer)


def workbook_from_response(response):
    return load_workbook(BytesIO(response.content))


def sheet_headers(workbook, sheet_name: str) -> list[str]:
    return [cell.value for cell in workbook[sheet_name][1]]


def test_export_products_generates_valid_xlsx_with_expected_columns(db_session: Session) -> None:
    workbook = workbook_from_buffer(ExportService(db_session).export_products())

    assert workbook.sheetnames == ["Products"]
    assert sheet_headers(workbook, "Products") == PRODUCT_EXPORT_COLUMNS
    assert workbook["Products"].freeze_panes == "A2"
    assert workbook["Products"]["A2"].value == 1
    assert workbook["Products"]["AD2"].value == 88


def test_export_recommended_products_filters_buy_and_review(db_session: Session) -> None:
    workbook = workbook_from_buffer(ExportService(db_session).export_recommended_products())
    sheet = workbook["Recommended Products"]
    statuses = [sheet[f"AE{row}"].value for row in range(2, sheet.max_row + 1)]

    assert statuses == ["buy", "review"]


def test_export_summary_report_contains_expected_sheets(db_session: Session) -> None:
    workbook = workbook_from_buffer(ExportService(db_session).export_summary_report())

    assert workbook.sheetnames == [
        "Summary",
        "Recommended Products",
        "Review Products",
        "Discarded Products",
        "Insufficient Data",
        "Suppliers",
        "Email Campaigns",
    ]
    summary = workbook["Summary"]
    assert summary["A2"].value == "Total products"
    assert summary["B2"].value == 3
    assert summary["A11"].value == "Top supplier by recommended products"
    assert summary["B11"].value == "Acme Supply (2)"


def test_export_suppliers_generates_expected_columns(db_session: Session) -> None:
    workbook = workbook_from_buffer(ExportService(db_session).export_suppliers())

    assert workbook.sheetnames == ["Suppliers"]
    assert sheet_headers(workbook, "Suppliers") == SUPPLIER_EXPORT_COLUMNS


@pytest.mark.parametrize(
    ("path", "filename", "sheet_name"),
    [
        ("/exports/products.xlsx", "products.xlsx", "Products"),
        ("/exports/recommended-products.xlsx", "recommended-products.xlsx", "Recommended Products"),
        ("/exports/suppliers.xlsx", "suppliers.xlsx", "Suppliers"),
        ("/exports/summary-report.xlsx", "summary-report.xlsx", "Summary"),
    ],
)
def test_export_endpoints_return_downloadable_xlsx(
    client: TestClient,
    path: str,
    filename: str,
    sheet_name: str,
) -> None:
    response = client.get(path)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"] == f'attachment; filename="{filename}"'
    assert sheet_name in workbook_from_response(response).sheetnames
